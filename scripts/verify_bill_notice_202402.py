# -*- coding: utf-8 -*-
"""verify_bill_notice_202402.py — 2024-02 催缴单全量对账(Excel 通知单 vs 引擎 bill_notice_line)。

左表 = 一期/二期两册全部「缴费通知单」sheet 的明细行(费项/表标识/用量/单价/金额),
      解析套路照抄 verify_tariff_rule_202402.py,并扩展抓取 容量/损耗/公摊/绿化水 行。
右表 = bill_notice_line JOIN bill_notice JOIN tenant (ym='2024-02')。
匹配 = 租户(名称+aliases,双向包含,全租户表,「X宿舍」家族折叠;匹配上但引擎无单→C类)
      → 费项类别 → 类内合计短路(±0.02 整类判中) → 段/金额贪心配对,容差 ±0.01。
只读对账:不改引擎、不写库。报告落 verify_bill_notice_202402_report.txt (UTF-8)。
"""
import io
import os
import re
import subprocess
import sys
import warnings
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP

warnings.filterwarnings("ignore")
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

BASE = r"C:\financial_dashboard\2025全年发生额、预算对比\2024年\2024年3月费用数据\2024年3月费用数据"
BOOKS = [("一期", os.path.join(BASE, "一期", "一期2024年2月水电费.xlsx"), "一期园区电", "一期园区水"),
         ("二期", os.path.join(BASE, "二期", "二期2024年2月水电费.xlsx"), "二期园区电", "二期园区水")]
YM = "2024-02"
NOTICE_MARK = "缴费通知单"
SKIP_SHEETS = {"厂房电", "厂房水", "火炬园（按总表差推算）"}
TOL = Decimal("0.01")
REPORT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "verify_bill_notice_202402_report.txt")

SEG_ALIAS = [("尖峰", "尖峰"), ("尖", "尖峰"), ("峰", "峰"), ("平", "平"), ("谷", "谷")]
SEG_NORM = {"尖峰": "sharp", "峰": "peak", "平": "flat", "谷": "valley"}
ENG_CAT = {"elec": "电", "mgmt_fee": "管理费", "capacity": "容量", "water": "水",
           "water_pipe": "管网", "share_elec_light": "公摊-路灯", "share_elec_elevator": "公摊-电梯",
           "share_elec_floor": "公摊-楼层消防", "share_elec_fire": "公摊-楼层消防",
           "share_green_water": "绿化水", "share_elec_loss": "损耗"}
ROW_CATS = ["电", "管理费", "容量", "水", "管网", "其他"]
# 结构性不可逐行:Excel 楼层公共+消防合一行/宿舍公摊按房间行,引擎按户按池分行 → 按户合计对账
SUM_CATS = ["公摊-路灯", "公摊-电梯", "公摊-楼层消防", "绿化水", "损耗"]

# §2.5 的 24 户例外 + 已知特殊户(用于不命中初步归因)
EXC24 = {"幸悦", "詹凯乔", "火炬园", "禹晨", "公交车站", "朱漫钳", "南一", "SENAN", "周兴",
         "威奈斯", "张丽莉", "张勤军", "沙力海", "芷泉", "袁华圣", "章肖艳", "李李商铺",
         "林锐辉", "星州", "永龙", "工程队宿舍", "粤海华创", "联塑精铟", "广联", "欧伟杰",
         "可盈", "开利暖通", "翔海", "张誉腾", "陈书谨", "可莱恩"}


def D(v):
    return None if v is None else Decimal(str(v))


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def r2(x):
    return x.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


# ────────────────────────────── Excel 解析(照抄 verify_tariff + 扩展) ──────────────────────────────
SRC_RE = re.compile(r"(宿舍电|宿舍水|一期园区电|一期园区水|二期园区电|二期园区水)")
PAREN_RE = re.compile(r"[（(][^）)]*[）)]")


def header_of(sv, r, maxc):
    cells = {}
    for c in range(1, maxc + 1):
        v = sv.cell(r, c).value
        if isinstance(v, str) and v.strip():
            cells.setdefault(v.strip(), c)
    if "宿舍" in cells and "租赁面积（㎡）" in cells:
        return ("dorm", cells)
    if "项目" in cells and ("上月行至" in cells or "上月抄表行至" in cells):
        return ("notice", cells)
    return None


def split_meter_label(item):
    s = str(item or "").replace("\n", "").strip()
    for sep in ("：", ":"):
        if sep in s:
            a, b = s.split(sep, 1)
            return a.strip(), b.strip()
    return "", s


def seg_of(label):
    s = PAREN_RE.sub("", str(label or "")).strip()
    for raw, norm in SEG_ALIAS:
        if s == raw or s.endswith(raw):
            return norm
    return None


def special_cat(lab):
    """明细行标签 → 费项类别(容量/损耗/公摊/绿化水/管网/其他);不含电/水(靠 src 判)。"""
    if lab.startswith(("基本用电费", "装机容量费", "新增基本用电费")):
        return "容量"
    if lab == "线路损耗":
        return "损耗"
    if lab.startswith(("路灯公摊", "路灯分摊")):
        return "公摊-路灯"
    if lab.startswith("电梯用电"):
        return "公摊-电梯"
    if lab in ("楼层公共、消防照明", "消防用电", "消防设施用电公摊"):
        return "公摊-楼层消防"
    if lab.startswith("绿化水公摊") or lab == "扣减收取已入驻企业绿化水费用":
        return "绿化水"
    if lab in ("电力管理费", "用电管理费"):
        return "管理费"
    if lab.startswith("水管网维护费") or lab == "用水维护费":
        return "管网"
    if lab in ("供水损耗", "扣减收取已入驻企业费用", "公共用电分摊", "公共用电东",
               "一至四车间高压用电分配", "公共用水费", "公共用水分摊"):
        return "其他"
    return None


def tagset(sv):
    return {str(sv.cell(r, 1).value).strip() for r in range(1, sv.max_row + 1)
            if isinstance(sv.cell(r, 1).value, str) and str(sv.cell(r, 1).value).strip()}


def parse_book(phase, wbf, wbv, park_e, park_w):
    tagsets = [("宿舍电", tagset(wbv["宿舍电"])), (park_e, tagset(wbv[park_e])),
               ("宿舍水", tagset(wbv["宿舍水"])), (park_w, tagset(wbv[park_w]))]
    rows, sheets = [], []
    for name in wbv.sheetnames:
        if name in SKIP_SHEETS:
            continue
        sv, sf = wbv[name], wbf[name]
        if NOTICE_MARK not in str(sv["B1"].value or ""):
            continue
        full = str(sv["B3"].value or "")  # 租户名称：xxx
        sheets.append((name, full.replace("租户名称：", "").split("|")[0].strip()))
        maxc = min(sv.max_column, 16)
        ctx, title, last_gkey, dorm_mgmt_p = None, str(sv["B1"].value or ""), None, None
        for r in range(1, sv.max_row + 1):
            b = sv.cell(r, 2).value
            if isinstance(b, str) and NOTICE_MARK in b:
                title = b
            h = header_of(sv, r, maxc)
            if h:
                ctx = h
                continue
            if ctx is None:
                continue
            kind, cm = ctx
            if kind == "notice":
                cp = cm["项目"]
                cprev = cm.get("上月行至") or cm.get("上月抄表行至")
                labels = [sv.cell(r, c).value for c in range(cp, cprev)]
                labels = [str(x).replace("\n", "").strip() for x in labels
                          if isinstance(x, str) and str(x).strip()]
                if not labels:
                    continue
                item = labels[-1]
                price = sv.cell(r, cm["单价"]).value if "单价" in cm else None
                qty = sv.cell(r, cm["实际用量"]).value if "实际用量" in cm else None
                amt = sv.cell(r, cm["金额"]).value if "金额" in cm else None
                tag = sv.cell(r, 1).value
                tag = tag.strip() if isinstance(tag, str) else None
                mprefix, lab = split_meter_label(item)
                sc = special_cat(lab)
                if sc:
                    if not num(amt):
                        continue
                    rows.append(dict(phase=phase, sheet=name, row=r, tag=tag or lab, cat=sc,
                                     seg=None, qty=D(qty) if num(qty) else None,
                                     price=D(price) if num(price) else None, amount=D(amt), lab=lab))
                    continue
                if not num(price):
                    continue
                fprev = sf.cell(r, cprev).value
                m = SRC_RE.search(str(fprev)) if fprev else None
                src = m.group(1) if m else None
                ccur = cm.get("本月行至")
                has_reading = num(sv.cell(r, cprev).value) and ccur and num(sv.cell(r, ccur).value)
                if src is None and tag and has_reading:
                    for sname, ts in tagsets:
                        if tag in ts:
                            src = sname
                            break
                if src is None:
                    continue
                seg = seg_of(lab)
                is_elec = src in ("宿舍电", "一期园区电", "二期园区电")
                is_maint = ("维护费" in title
                            or str(labels[0] if len(labels) > 1 else "").startswith(
                                ("用电维护费", "水电维护费", "用水维护费")))
                gkey = (name, tag, mprefix)
                if is_elec:
                    # 双票结构:维护费单把电表行按 0.16/0.32 重列 → 归入管理费类,按表聚合
                    cat = "管理费" if is_maint else "电"
                    rows.append(dict(phase=phase, sheet=name, row=r, tag=tag, cat=cat,
                                     seg=SEG_NORM.get(seg), qty=D(qty) if num(qty) else None,
                                     price=D(price), amount=D(amt) if num(amt) else Decimal(0),
                                     lab=lab, gkey=gkey, maint=is_maint))
                else:
                    cat = "管网" if is_maint else "水"
                    rows.append(dict(phase=phase, sheet=name, row=r, tag=tag, cat=cat,
                                     seg=None, qty=D(qty) if num(qty) else None, price=D(price),
                                     amount=D(amt) if num(amt) else Decimal(0), lab=lab))
            else:  # 宿舍段分间块
                room = sv.cell(r, cm["宿舍"]).value
                citem = cm.get("项目") or cm["租赁面积（㎡）"] + 1
                item = sv.cell(r, citem).value
                item = item.strip() if isinstance(item, str) else ""
                amt = sv.cell(r, cm["金额"]).value if "金额" in cm else None
                if item in ("路灯分摊", "绿化水公摊"):
                    if num(amt):
                        rows.append(dict(phase=phase, sheet=name, row=r, tag=f"宿舍{room or ''}",
                                         cat="公摊-路灯" if item == "路灯分摊" else "绿化水",
                                         seg=None, qty=None, price=None, amount=D(amt), lab=item))
                    continue
                if room is None or item not in ("电表", "水表", "新水表"):
                    continue
                pcol = cm.get("基准电价") or cm.get("单价")
                price = sv.cell(r, pcol).value if pcol else None
                if not num(price):
                    continue
                qty = sv.cell(r, cm["实际用量"]).value if "实际用量" in cm else None
                tag = f"宿舍{room}"
                if item == "电表":
                    q = D(qty) if num(qty) else Decimal(0)
                    # Excel 宿舍行金额 = 度数×(基准价+管理费) 合并列;拆成 电 + 管理费 两虚拟行对齐引擎
                    rows.append(dict(phase=phase, sheet=name, row=r, tag=tag, cat="电", seg=None,
                                     qty=q, price=D(price), amount=r2(q * D(price)), lab="宿舍电表",
                                     dorm=True))
                    if "电力管理费" in cm:
                        mp = sv.cell(r, cm["电力管理费"]).value
                        if num(mp):
                            rows.append(dict(phase=phase, sheet=name, row=r, tag=tag, cat="管理费",
                                             seg=None, qty=q, price=D(mp), amount=r2(q * D(mp)),
                                             lab="宿舍电力管理费", dorm=True))
                else:
                    rows.append(dict(phase=phase, sheet=name, row=r, tag=tag, cat="水", seg=None,
                                     qty=D(qty) if num(qty) else None, price=D(price),
                                     amount=D(amt) if num(amt) else Decimal(0), lab="宿舍水表",
                                     dorm=True))
    return rows, sheets


def fold_maint(rows):
    """维护费单重列的电表分段行(管理费类)按表聚合成一行,对齐引擎的一表一条 mgmt_fee。"""
    out, buf = [], defaultdict(list)
    for r in rows:
        if r["cat"] == "管理费" and r.get("maint") and r.get("gkey"):
            buf[(r["sheet"], r["gkey"])].append(r)
        else:
            out.append(r)
    for (_, gk), lst in buf.items():
        base = dict(lst[0])
        base["amount"] = sum(x["amount"] for x in lst)
        base["qty"] = sum(x["qty"] for x in lst if x["qty"] is not None)
        base["seg"] = None
        base["lab"] = "维护费单聚合"
        out.append(base)
    return out


# ────────────────────────────── 引擎侧 ──────────────────────────────
def mysql(sql):
    p = subprocess.run(["docker", "exec", "demo3-mysql", "mysql", "-uroot", "-proot",
                        "--default-character-set=utf8mb4", "-B", "-N", "park_demo3", "-e", sql],
                       capture_output=True)
    return p.stdout.decode("utf-8", errors="replace").splitlines()


def load_engine():
    sql = ("SELECT n.tenant_id, t.company_name, IFNULL(t.aliases,''), n.notice_kind, l.fee_key, "
           "IFNULL(l.seg,''), IFNULL(l.meter_label,''), IFNULL(l.qty,''), l.amount, "
           "REPLACE(REPLACE(IFNULL(l.note,''),'\\t',' '),'\\n',' ') "
           "FROM bill_notice_line l JOIN bill_notice n ON n.id=l.notice_id "
           "JOIN tenant t ON t.id=n.tenant_id WHERE n.ym='%s' ORDER BY n.id,l.line_no" % YM)
    lines = []
    tenants = {}
    for ln in mysql(sql):
        f = ln.split("\t")
        if len(f) < 10:
            continue
        tid = int(f[0])
        tenants[tid] = (f[1], f[2])
        lines.append(dict(tid=tid, name=f[1], kind=f[3], fee_key=f[4], seg=f[5] or None,
                          meter=f[6], qty=D(f[7]) if f[7] not in ("", "NULL") else None,
                          amount=D(f[8]), note=f[9], cat=ENG_CAT.get(f[4], "其他")))
    # 全租户表(不只出单户):sheet 匹配用;匹配上但无 2024-02 单 → C类
    all_tenants = {}
    for ln in mysql("SELECT id, company_name, IFNULL(aliases,'') FROM tenant"):
        f = ln.split("\t")
        if len(f) >= 3:
            all_tenants[int(f[0])] = (f[1], f[2])
    return lines, tenants, all_tenants


def build_name_map(tenants):
    m = defaultdict(set)
    for tid, (name, aliases) in tenants.items():
        m[name].add(tid)
        for a in re.split(r"[,，、;；/|\s]+", aliases):
            if a:
                m[a].add(tid)
    return m


def build_canon(all_tenants, name_map):
    """家族折叠:「X宿舍」租户折到「X」租户(龙为宿舍→龙为)。"""
    canon = {}
    for tid, (name, _) in all_tenants.items():
        if name.endswith("宿舍"):
            tids = name_map.get(name[:-2], set())
            if len(tids) == 1:
                canon[tid] = next(iter(tids))
    return canon


INNER_PAREN_RE = re.compile(r"[（(]([^）)]+)[）)]")


def resolve_sheet(sheet, full, name_map, canon):
    if sheet in name_map:
        cands = {canon.get(t, t) for t in name_map[sheet]}
        if len(cands) == 1:
            return next(iter(cands))
    # 双向包含:拆短名(去括号)+括号全名两路;短名⊆company_name 或 company_name⊆全名
    names = {sheet, PAREN_RE.sub("", sheet).strip()}
    m = INNER_PAREN_RE.search(sheet)
    if m:
        names.add(m.group(1).strip())
    # 家族折叠:「X宿舍」sheet 折成「X」再试
    names |= {n.replace("宿舍", "").strip() for n in list(names) if "宿舍" in n}
    names.discard("")
    cands = set()
    for nm, tids in name_map.items():
        if not nm:
            continue
        for s in names:
            if nm in s or s in nm or (full and (nm in full or full in nm)):
                cands |= tids
                break
    cands = {canon.get(t, t) for t in cands}
    return next(iter(cands)) if len(cands) == 1 else None


# ────────────────────────────── 匹配 ──────────────────────────────
def greedy_match(ex, en):
    """同(租户,类别)桶内:先 段+金额 全等,再 仅金额(±0.01)。返回配对数与双侧剩余。"""
    used_e, used_g = set(), set()
    pairs = 0
    for strict in (True, False):
        for i, xr in enumerate(ex):
            if i in used_e:
                continue
            for j, gr in enumerate(en):
                if j in used_g:
                    continue
                if strict and xr["seg"] != gr["seg"]:
                    continue
                if abs(xr["amount"] - gr["amount"]) <= TOL:
                    used_e.add(i)
                    used_g.add(j)
                    pairs += 1
                    break
    return pairs, [x for i, x in enumerate(ex) if i not in used_e], \
        [g for j, g in enumerate(en) if j not in used_g]


NOTE_AMT_RE = re.compile(r"金额口径=.*?=(-?[\d.]+)\(")
NOTE_DORM_RE = re.compile(r"宿舍段损耗=")


def attribute(tenant_disp, cat, ex_rows, en_rows, sum_ok, side):
    if any(k in tenant_disp for k in EXC24):
        return "例外未录待核(§2.5 24户)"
    if side == "excel" and not en_rows:
        return "整类缺失→绑定/派生缺口"
    if side == "engine" and not ex_rows:
        return "引擎多派生(Excel该户无此类)"
    if sum_ok:
        return "粒度差异(类内合计相等)"
    if cat == "电" and any(r.get("dorm") for r in (ex_rows or [])):
        return "宿舍段口径差?"
    return "待核(引擎缺陷?)"


def main():
    ex_rows, sheet_list = [], []
    for phase, path, pe, pw in BOOKS:
        wbf = openpyxl.load_workbook(path, data_only=False)
        wbv = openpyxl.load_workbook(path, data_only=True)
        rows, sheets = parse_book(phase, wbf, wbv, pe, pw)
        ex_rows += rows
        sheet_list += [(phase, s, f) for s, f in sheets]
    ex_rows = fold_maint(ex_rows)

    en_lines, tenants, all_tenants = load_engine()
    name_map_all = build_name_map(all_tenants)
    canon = build_canon(all_tenants, name_map_all)
    for g in en_lines:  # 引擎侧同步折叠(龙为宿舍单并入龙为)
        g["tid"] = canon.get(g["tid"], g["tid"])
    billed = {g["tid"] for g in en_lines}
    name_map_billed = build_name_map({t: v for t, v in all_tenants.items()
                                      if canon.get(t, t) in billed})

    sheet_tid, unresolved = {}, []
    for phase, s, f in sheet_list:
        # 先在出单户里试(保持原精度),再放宽到全租户表
        tid = resolve_sheet(s, f, name_map_billed, canon)
        if tid is None:
            tid = resolve_sheet(s, f, name_map_all, canon)
        if tid is None:
            unresolved.append((phase, s, f))
        else:
            sheet_tid[(phase, s)] = tid

    # 桶化:key=(tid或伪键, 类别);匹配到租户但引擎无单 → C类,不进不命中池
    exb, enb = defaultdict(list), defaultdict(list)
    c_rows = defaultdict(lambda: [0, Decimal(0)])  # tid -> [行数, 金额]
    for r in ex_rows:
        tid = sheet_tid.get((r["phase"], r["sheet"]))
        r["tkey"] = tid if tid is not None else "SHEET:%s·%s" % (r["phase"], r["sheet"])
        if tid is not None and tid not in billed:
            r["cskip"] = True
            c_rows[tid][0] += 1
            c_rows[tid][1] += r["amount"]
            continue
        exb[(r["tkey"], r["cat"])].append(r)
    for g in en_lines:
        enb[(g["tid"], g["cat"])].append(g)

    def disp(tkey):
        if isinstance(tkey, int):
            return all_tenants[tkey][0]
        return tkey

    out = io.StringIO()
    W = out.write
    W("══════════ 催缴单全量对账 %s (Excel 通知单 vs 引擎 bill_notice_line) ══════════\n" % YM)
    W("源册: %s\n      %s\n" % (BOOKS[0][1], BOOKS[1][1]))
    W("通知单 sheet: %d 张(其中未匹配到租户 %d 张); 引擎: %d 行 / %d 租户\n" %
      (len(sheet_list), len(unresolved), len(en_lines), len(tenants)))
    if unresolved:
        W("未匹配 sheet: %s\n" % ", ".join("%s·%s(%s)" % u for u in unresolved))
    if c_rows:
        W("C类·匹配到租户但引擎无 %s 单 (%d 户,%d 行,不进不命中池): %s\n" %
          (YM, len(c_rows), sum(v[0] for v in c_rows.values()),
           ", ".join("%s(%d行/%s元)" % (disp(t), v[0], r2(v[1])) for t, v in
                     sorted(c_rows.items(), key=lambda kv: -kv[1][1]))))
    W("\n")

    # ── 逐行类:类内合计短路 → 桶内贪心配对 ──
    stat = defaultdict(lambda: [0, 0, 0, 0])  # cat -> [ex_n, en_n, hit_ex, hit_en]
    miss_ex, miss_en, diffs = [], [], []
    zero_ex, zero_en = defaultdict(int), defaultdict(int)  # 0 金额孤行只计数不列明细
    SC_TOL = Decimal("0.02")
    sc_buckets = []  # 合计短路且行数不等(粒度差)的桶
    for key in {k for k in (set(exb) | set(enb)) if k[1] in ROW_CATS}:
        tkey, cat = key
        ex, en = exb.get(key, []), enb.get(key, [])
        # 类内合计短路:两侧都有且合计相等(±0.02) → 整类判中,不进不命中池
        if ex and en and abs(sum(x["amount"] for x in ex)
                             - sum(g["amount"] for g in en)) <= SC_TOL:
            stat[cat][0] += len(ex)
            stat[cat][1] += len(en)
            stat[cat][2] += len(ex)
            stat[cat][3] += len(en)
            if len(ex) != len(en):
                sc_buckets.append((disp(tkey), cat, len(ex), len(en),
                                   sum(x["amount"] for x in ex)))
            continue
        hit, rem_ex, rem_en = greedy_match(ex, en)
        sum_ok = abs(sum(x["amount"] for x in ex) - sum(g["amount"] for g in en)) \
            <= TOL * max(1, len(ex) + len(en))
        stat[cat][0] += len(ex)
        stat[cat][1] += len(en)
        stat[cat][2] += hit
        stat[cat][3] += hit
        # 剩余两侧行再按 段同→顺序 结对,展示 Excel值 vs 引擎值(疑价差/量差)
        for strict in (True, False):
            for x in list(rem_ex):
                if x["amount"] == 0:
                    continue
                for g in list(rem_en):
                    if g["amount"] == 0:
                        continue
                    if strict and x["seg"] != g["seg"]:
                        continue
                    diffs.append((disp(tkey), cat, x, g))
                    rem_ex.remove(x)
                    rem_en.remove(g)
                    break
        for x in rem_ex:
            if x["amount"] == 0:
                zero_ex[cat] += 1
            else:
                miss_ex.append((disp(tkey), cat, x, en, sum_ok))
        for g in rem_en:
            if g["amount"] == 0:
                zero_en[cat] += 1
            else:
                miss_en.append((disp(tkey), cat, g, ex, sum_ok))

    # ── 合计类:按户合计对账 ──
    sum_stat = {}          # cat -> [n_tenant, hit]
    sum_miss = defaultdict(list)  # cat -> [(tenant, ex_total, en_total)]
    TOLS = Decimal("0.05")
    for cat in SUM_CATS:
        ex_t, en_t = defaultdict(Decimal), defaultdict(Decimal)
        rows_e = rows_g = 0
        for (tkey, c), lst in exb.items():
            if c == cat:
                ex_t[tkey] += sum(x["amount"] for x in lst)
                rows_e += len(lst)
        for (tid, c), lst in enb.items():
            if c == cat:
                en_t[tid] += sum(g["amount"] for g in lst)
                rows_g += len(lst)
        tks = set(ex_t) | set(en_t)
        hit = 0
        for tk in tks:
            a, b = ex_t.get(tk, Decimal(0)), en_t.get(tk, Decimal(0))
            if abs(a - b) <= TOLS:
                hit += 1
            else:
                sum_miss[cat].append((disp(tk), a, b))
        sum_stat[cat] = [len(tks), hit, rows_e, rows_g]

    W("── 总览 A:逐行类(类内合计短路±0.02 → 桶内 段+金额 贪心配对,容差±0.01) ──\n")
    W("%-12s %8s %8s %10s %10s %10s %14s\n" %
      ("类别", "Excel行", "引擎行", "命中(E/引擎)", "Excel命中率", "引擎命中率", "0元孤行(E/引擎)"))
    te = tg = the = thg = 0
    for cat in ROW_CATS:
        e, g, he, hg = stat[cat]
        if e == g == 0:
            continue
        te += e
        tg += g
        the += he
        thg += hg
        W("%-12s %8d %8d %10s %9s %9s %10s\n" % (cat, e, g, "%d/%d" % (he, hg),
          "%.1f%%" % (he / e * 100) if e else "-", "%.1f%%" % (hg / g * 100) if g else "-",
          "%d/%d" % (zero_ex[cat], zero_en[cat])))
    W("%-12s %8d %8d %10s %9s %9s\n" % ("合计", te, tg, "%d/%d" % (the, thg),
      "%.1f%%" % (the / te * 100) if te else "-", "%.1f%%" % (thg / tg * 100) if tg else "-"))
    if sc_buckets:
        W("类内合计短路·行数不等(粒度差,整类判中): %s\n" %
          "; ".join("%s·%s %d行↔%d行=%s元" % b for b in
                    sorted(sc_buckets, key=lambda t: (t[1], t[0]))))
    W("\n")

    W("── 总览 B:合计类(按户合计对账,容差±0.05;Excel合并行/宿舍分间行 vs 引擎分池行,不可逐行) ──\n")
    W("%-12s %8s %8s %8s %8s %10s\n" % ("类别", "Excel行", "引擎行", "涉及户数", "户命中", "户命中率"))
    for cat in SUM_CATS:
        n, h, re_, rg_ = sum_stat[cat]
        W("%-12s %8d %8d %8d %8d %9s\n" %
          (cat, re_, rg_, n, h, "%.1f%%" % (h / n * 100) if n else "-"))
    W("\n")

    # ── 损耗专项 ──
    W("── 损耗行专项:Excel 金额 vs 引擎 amount(度数口径) vs 引擎 note 金额口径 ──\n")
    loss_ex = defaultdict(Decimal)
    for r in ex_rows:
        if r["cat"] == "损耗" and not r.get("cskip"):
            loss_ex[r["tkey"]] += r["amount"]
    loss_en = defaultdict(lambda: [Decimal(0), Decimal(0), Decimal(0)])  # amt, note_amt, dorm_amt
    for g in en_lines:
        if g["fee_key"] != "share_elec_loss":
            continue
        if NOTE_DORM_RE.search(g["note"]):
            loss_en[g["tid"]][2] += g["amount"]  # 宿舍段:amount 本身即金额口径,Excel 无对应行
            continue
        loss_en[g["tid"]][0] += g["amount"]
        m = NOTE_AMT_RE.search(g["note"])
        loss_en[g["tid"]][1] += Decimal(m.group(1)) if m else Decimal(0)
    tol_l = Decimal("0.02")
    w_amt = w_note = w_none = 0
    W("%-16s %10s %12s %12s  %s\n" % ("租户", "Excel损耗", "引擎amount", "note金额口径", "命中"))
    all_tids = sorted(set(loss_ex) | set(loss_en), key=str)
    rows_out = []
    for tk in all_tids:
        exv = loss_ex.get(tk, Decimal(0))
        amt, namt, dorm = loss_en.get(tk, [Decimal(0), Decimal(0), Decimal(0)]) if isinstance(tk, int) else [Decimal(0)] * 3
        hit_a = abs(exv - amt) <= tol_l
        hit_n = abs(exv - namt) <= tol_l
        tag = "两口径同" if hit_a and hit_n else ("度数口径" if hit_a else ("金额口径" if hit_n else "均不中"))
        if hit_n:
            w_note += 1
        if hit_a:
            w_amt += 1
        if not hit_a and not hit_n:
            w_none += 1
        rows_out.append("%-16s %10s %12s %12s  %s%s" %
                        (str(disp(tk))[:15], exv, amt, namt, tag,
                         " (另宿舍段损耗 %s,Excel无此行)" % dorm if dorm else ""))
    W("\n".join(rows_out) + "\n")
    W("口径统计(±0.02): 金额口径命中 %d 户 / 度数口径命中 %d 户 / 均不中 %d 户 (共 %d 户)\n" %
      (w_note, w_amt, w_none, len(all_tids)))
    n_amt = sum(1 for tk in all_tids
                if abs(loss_ex.get(tk, Decimal(0)) - (loss_en.get(tk, [Decimal(0)] * 3)[0]
                       if isinstance(tk, int) else Decimal(0))) <= 1)
    n_note = sum(1 for tk in all_tids
                 if abs(loss_ex.get(tk, Decimal(0)) - (loss_en.get(tk, [Decimal(0)] * 3)[1]
                        if isinstance(tk, int) else Decimal(0))) <= 1)
    W("口径倾向(±1元): 金额口径近中 %d 户 / 度数口径近中 %d 户\n" % (n_note, n_amt))
    dorm_total = sum(v[2] for v in loss_en.values())
    W("宿舍段损耗(引擎有/Excel通知单无独立行): 合计 %s 元\n\n" % dorm_total)

    # ── 不命中清单 ──
    def sum_attr(tn, a, b):
        if any(k in tn for k in EXC24):
            return "例外未录待核(§2.5 24户)"
        if b == 0:
            return "整类缺失→绑定/派生缺口"
        if a == 0:
            return "引擎多派生(Excel无此行)"
        if abs(a - b) < 1:
            return "尾差<1元(舍入/基数微差)"
        return "待核(池份额/引擎缺陷?)"

    def diff_attr(tn, x, g):
        if any(k in tn for k in EXC24):
            return "例外未录待核(§2.5 24户)"
        if x.get("qty") is not None and g["qty"] is not None and abs(x["qty"] - g["qty"]) <= Decimal("0.01"):
            return "量同价异(单价口径?)"
        if x.get("price") is not None and g["qty"] not in (None, 0) \
                and abs(x["amount"] / (x["qty"] or 1) - g["amount"] / g["qty"]) <= Decimal("0.02"):
            return "价同量异(读数/绑定?)"
        return "待核(引擎缺陷?)"

    W("── 不命中 · 两侧都有但金额不等 (%d 对,同桶内剩余行结对) ──\n" % len(diffs))
    W("%-16s %-10s %-18s %10s %10s  %s\n" % ("租户", "类别", "标签/表", "Excel值", "引擎值", "初步归因"))
    for tn, cat, x, g in sorted(diffs, key=lambda t: (t[1], t[0])):
        W("%-16s %-10s %-18s %10s %10s  %s\n" %
          (tn[:15], cat, ("%s/%s" % (x.get("tag") or "", x.get("lab") or ""))[:18],
           x["amount"], g["amount"], diff_attr(tn, x, g)))
    W("\n")

    sm_ex = [(tn, cat, a, b) for cat in SUM_CATS for tn, a, b in sum_miss[cat] if a > b]
    sm_en = [(tn, cat, a, b) for cat in SUM_CATS for tn, a, b in sum_miss[cat] if b > a]
    W("── 不命中 · Excel 有而引擎无 (逐行类 %d 行 + 合计类 %d 户) ──\n" % (len(miss_ex), len(sm_ex)))
    W("%-16s %-10s %-18s %10s  %s\n" % ("租户", "类别", "标签/表", "Excel金额", "初步归因"))
    for tn, cat, x, en, sum_ok in sorted(miss_ex, key=lambda t: (t[1], t[0])):
        W("%-16s %-10s %-18s %10s  %s\n" %
          (tn[:15], cat, ("%s/%s" % (x.get("tag") or "", x.get("lab") or ""))[:18],
           x["amount"], attribute(tn, cat, [x], en, sum_ok, "excel")))
    for tn, cat, a, b in sorted(sm_ex, key=lambda t: (t[1], t[0])):
        W("%-16s %-10s %-18s %10s  %s\n" %
          (tn[:15], cat, "户合计 引擎=%s" % b, a, sum_attr(tn, a, b)))
    W("\n── 不命中 · 引擎有而 Excel 无 (逐行类 %d 行 + 合计类 %d 户) ──\n" % (len(miss_en), len(sm_en)))
    W("%-16s %-10s %-18s %10s  %s\n" % ("租户", "类别", "fee_key/seg", "引擎金额", "初步归因"))
    for tn, cat, g, ex, sum_ok in sorted(miss_en, key=lambda t: (t[1], t[0])):
        W("%-16s %-10s %-18s %10s  %s\n" %
          (tn[:15], cat, ("%s/%s" % (g["fee_key"], g["seg"] or ""))[:18],
           g["amount"], attribute(tn, cat, ex, [g], sum_ok, "engine")))
    for tn, cat, a, b in sorted(sm_en, key=lambda t: (t[1], t[0])):
        W("%-16s %-10s %-18s %10s  %s\n" %
          (tn[:15], cat, "户合计 Excel=%s" % a, b, sum_attr(tn, a, b)))

    text = out.getvalue()
    with open(REPORT, "w", encoding="utf-8") as f:
        f.write(text)
    print(text)
    print("报告已落:", REPORT)


if __name__ == "__main__":
    main()
