# -*- coding: utf-8 -*-
"""extract.py — 租金通知单源册逐户抽成结构化数据(合同核对用),带合计自检。

    python scripts/contract-recon/extract.py            → scripts/contract-recon/out/

产物(out/):
  <册id>.json           每本租金册一份:sheet → 通知单(抬头/月份/类型/逐行/合计/落款)+ 嵌图
  all.json              总表:册清单(zone/月/sheet 数/重复关系)、各 zone 户数、全部非重复册的通知单拍平、自检失败
  images/<册id>/<sheet>_<n>.png   嵌图导出(纸质合同截图);重复册不导出,看 duplicateOf 那本;
                        同一张图(sha1 相同)各月反复嵌,只落一次盘,其余是硬链接
  contract-summary-2024-03.json   园区租户租金合同明细汇总(2024年3月).xlsx「汇总」
  lease-status-<日期>.json        园区面积、租金、管理租赁情况表(最新一版,全部 sheet)
  utility-area.json     水电费册(pool-recon MONTHS)各户通知单的 路灯公摊/绿化水公摊 行面积与抬头位置

版式按内容认,不按行号:通知单按 B 列含「通知单」的标题行切开;表头是含「收费项目/物业名称/位置」的行,
列按表头文字归到规范字段(COLS);表体到「合计」行为止。数取单元格缓存值(data_only)。
「收费项目|物业名称」表头多数和内容对调,按内容(哪列像房号/楼座)认 location,换了记 headerSwapped。
自检:每张通知单「合计」= Σ应收金额(容差 0.01),不等的列在 all.json selfcheck 里,退出码 1。
"""
import glob
import hashlib
import importlib.util
import json
import os
import posixpath
import re
import sys
import warnings
import zipfile
from concurrent.futures import ProcessPoolExecutor
from datetime import datetime
from decimal import Decimal
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter as L

warnings.filterwarnings("ignore")

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out")
_spec = importlib.util.spec_from_file_location("pool_extract", os.path.join(HERE, "..", "pool-recon", "extract.py"))
PX = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(PX)
SRC, D, num, ym_of = PX.SRC, PX.D, PX.num, PX.ym_of      # 源册根目录 / 数值工具 / 水电费册 MONTHS 都从 pool-recon 取

RENT = re.compile(r"^(一期|二期)(\d{4})年(\d{1,2})月租金")
SUMMARY = r"2024年\2024年3月费用数据\园区租户租金合同明细汇总(2024年3月).xlsx"
LEASE_GLOB = r"2023年\园区面积、租金、管理租赁情况表（*）.xlsx"
TOL = Decimal("0.01")
FEE_LIKE = r"(租金|费|税|押金|保证金)$"
LOC_LIKE = r"[座室楼层栋房间号]|车间|空地|通道|饭堂"
MAX_COL = 30

# 表头文字(去空白、全角括号转半角、去 ㎡)→ 规范字段。按顺序第一个命中的算。
COLS = [
    (r"^物业名称", "location"), (r"^位置", "position"), (r"^收费项目", "fee"), (r"^项目", "item"),
    (r"^期间", "period"), (r"^房号", "room"), (r"^房数", "rooms"), (r"^系数", "coef"), (r"^税率", "taxRate"),
    (r"^空地面积", "landArea"), (r"^分摊面积单价|^分摊单价", "shareUnitPrice"),
    (r"^(建筑)?面积单价|^单价/元\(含税\)|^单价\(元/\)$|^单价$|^单价/元$|^单价\(\d+%部分\)", "unitPrice"),
    (r"^单价/元\(不含税\)", "unitPriceExTax"),
    (r"月单价|^单价\(元\)$|^单价\(元/月\)|^元/月|^月金额", "monthly"),
    (r"^建筑面积", "buildArea"), (r"^分摊面积", "shareArea"), (r"^(套内|户内)面积", "innerArea"),
    (r"^计费面积", "billArea"), (r"^面积", "area"), (r"金额", "amount"), (r"^备注", "remark"),
]


def s(v):
    return re.sub(r"\s", "", str(v)) if v is not None else ""


def val(v):
    """JSON 出口:数 → num() 字符串(保精度);日期 → ISO;其余原样(去首尾空白)。"""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return num(v)
    return v.strip() if isinstance(v, str) else v


def label(v):
    return s(v).replace("（", "(").replace("）", ")").replace("㎡", "").replace("m²", "")


def map_header(g, hr):
    """表头行 → [(列, 规范字段, 原文)],到「备注」列为止;备注右边的格记 extra。"""
    cols, used, extra = [], set(), []
    labels = sorted((c, g[(r, c)]) for (r, c) in g if r == hr)
    has_loc = any(label(v).startswith("物业名称") for _, v in labels)
    dup_build = sum(1 for _, v in labels if label(v).startswith("建筑面积") and "单价" not in label(v)) > 1
    done = False
    for c, v in labels:
        if done:
            extra.append(val(v))
            continue
        t = label(v)
        key = next((k for p, k in COLS if re.search(p, t)), None) or f"col{L(c)}"
        if key == "buildArea" and dup_build and "landArea" not in used:
            key = "landArea"       # 「建筑面积|建筑面积」版式:第一列在模板里是空地面积位,表头抄错
        if key == "position" and not has_loc:
            key = "location"       # 「位置|收费项目|…」版式:没有物业名称列,位置就是物业名称
        while key in used:
            key += "_2"
        used.add(key)
        cols.append((c, key, str(v).strip()))
        done = key == "remark"
    return cols, extra


def is_loc(v):
    t = s(v)
    return bool(re.search(LOC_LIKE, t)) and not re.search(FEE_LIKE, t)


def is_title(t):
    return "通知单" in t and len(t) < 60 and "签收" not in t and "接收" not in t


def parse_rent_sheet(g):
    """一个 sheet 的缓存值网格 {(r,c): v} → 通知单列表(按 B 列标题行切)。"""
    titles = sorted(r for (r, c), v in g.items() if c == 2 and isinstance(v, str) and is_title(s(v)))
    max_r = max((r for r, _ in g), default=0)
    notices = []
    for i, tr in enumerate(titles):
        end = titles[i + 1] if i + 1 < len(titles) else max_r + 1
        seg = sorted((r, c, v) for (r, c), v in g.items() if tr <= r < end)
        title = str(g[(tr, 2)]).strip()
        n = {"titleRow": tr, "title": title, "kind": re.sub(r"^[\d年月日\-—~至起.\s]+", "", s(title)),
             "ym": ym_of(title), "titleNote": [val(v) for r, c, v in seg if r == tr and c != 2]}
        m = re.search(r"(\d{4})年(\d{1,2})月?(?:\d{1,2}日)?(?:起)?[-—~至](?:(\d{4})年)?(\d{1,2})月", s(title))
        if m:
            n["ymTo"] = f"{m.group(3) or m.group(1)}-{int(m.group(4)):02d}"
        texts = [(r, c, str(v)) for r, c, v in seg if isinstance(v, str)]
        to = next(((r, t) for r, c, t in texts if re.match(r"\s*(尊敬的|Dear)", t, re.I)), None)
        n["addressee"] = re.sub(r"^\s*(尊敬的|Dear)\s*|[\s:：,，]+$", "", to[1], flags=re.I) if to else None
        blob = " ".join(t for _, _, t in texts)
        due = re.search(r"于(\d{4}年\d{1,2}月\d{1,2}日)前", blob)
        payee = re.search(r"户名[:：]\s*(\S+)", blob)
        n["dueDate"], n["payee"] = due and due.group(1), payee and payee.group(1)
        hr = next((r for r, c, t in texts if s(t) in ("收费项目", "物业名称", "位置")), None)
        n["headerRow"], rows, total, trow = hr, [], None, None
        if hr is None:
            n["problem"] = "无表头(找不到「收费项目/物业名称/位置」)"
        else:
            cols, n["headerExtra"] = map_header(g, hr)
            n["header"] = [h for _, _, h in cols]
            key_of = {c: k for c, k, _ in cols}
            for c in range(2, max(key_of) + 1):     # 表头空着的列(如「位置|…|金额」版式的 C/D)也收,记 colX
                key_of.setdefault(c, f"col{L(c)}")
            by_key = {k: c for c, k in key_of.items()}
            if "location" in by_key and "fee" in by_key:
                # 「收费项目|物业名称」表头多数和内容对调(B 列实为位置),按内容认:哪列位置样的值多,哪列就是 location
                t_end = next((r for r in range(hr + 1, end) if s(g.get((r, 2))).startswith(("合计", "总计"))), end)

                def loc_like(c):       # 只数表体:4 月册的「应收费用」矩阵在通知单下方
                    return sum(1 for r in range(hr + 1, t_end) if is_loc(g.get((r, c))))
                if loc_like(by_key["fee"]) > loc_like(by_key["location"]):
                    key_of[by_key["location"]], key_of[by_key["fee"]] = "fee", "location"
                    n["headerSwapped"] = True
            amt_c = next((c for c, k, _ in cols if k == "amount"), None)
            label_cs = {2} | {c for c, k, _ in cols if k in ("location", "fee", "item")}
            loc = None
            for r in range(hr + 1, end):
                cells = {c: g.get((r, c)) for c in key_of}
                if any(s(g.get((r, c))).startswith(("合计", "总计")) for c in label_cs):
                    trow = r
                    total = D(g.get((r, amt_c))) if amt_c else None
                    if total is None:
                        total = next((D(v) for (rr, c), v in sorted(g.items()) if rr == r and D(v) is not None), None)
                    break
                if any(k in s(v) for v in cells.values() if isinstance(v, str) for k in ("制表", "签收")):
                    break
                if all(v is None or s(v) == "" for v in cells.values()):
                    continue
                row = {"row": r}
                for c, k in key_of.items():
                    if cells[c] is not None and s(cells[c]) != "":
                        row[k] = val(cells[c])
                if ("fee" not in row and "location" in row and re.search(FEE_LIKE, s(row["location"]))
                        and not is_loc(row["location"])):
                    row["fee"], row["feeFromLocation"] = row.pop("location"), True    # 费目名写在 B:C 合并格里
                if "location" in row:
                    loc = row["location"]
                elif loc is not None:
                    row["location"], row["locFilled"] = loc, True     # 物业名称合并格/留空:沿用上一行
                rows.append(row)
        n["rows"] = rows
        n["totalRow"], n["total"] = trow, num(total)
        ssum = sum((D(x.get("amount")) or Decimal(0) for x in rows), Decimal(0))
        n["sumAmount"] = num(ssum)
        n["ok"] = total is not None and abs(total - ssum) <= TOL
        # 签收企业 / 落款:签收栏下方第一个含「公司」的格是落款,其后第一个日期是落款日期
        after = [(r, c, v) for r, c, v in seg if trow is None or r > trow]
        signee = next((g.get((r, c2)) for r, c, v in after if s(v) == "签收企业"
                       for c2 in range(c + 1, c + 4) if g.get((r, c2)) is not None), None)
        n["signee"] = val(signee)
        sd = next((r for r, c, v in after if s(v) == "签收日期"), None)
        n["signer"] = next((val(v) for r, c, v in after if (sd is None or r > sd) and isinstance(v, str) and "公司" in v
                            and "签收" not in v), None)
        n["signDate"] = next((val(v) for r, c, v in after if (sd is None or r > sd) and isinstance(v, datetime)), None)
        notices.append(n)
    return notices


# ───────────────────────────── 读册 ─────────────────────────────
def read_grids(path, max_col=MAX_COL):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    grids = {}
    for name in wb.sheetnames:
        g = {}
        for r, row in enumerate(wb[name].iter_rows(max_col=max_col, values_only=True), 1):
            for c, v in enumerate(row, 1):
                if v is not None:
                    g[(r, c)] = v
        grids[name] = g
    return wb.sheetnames, grids


NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def _rels(z, part):
    rp = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if rp not in z.namelist():
        return {}
    out = {}
    for e in ET.fromstring(z.read(rp)):
        t = e.get("Target")
        t = t[1:] if t.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(part), t))
        out[e.get("Id")] = (e.get("Type", "").rsplit("/", 1)[-1], t)
    return out


def sheet_media(path):
    """直接读 xlsx 包:sheet 名 → {state, images:[{media, row, col}]}(锚点左上角格)。"""
    z = zipfile.ZipFile(path)
    wr = _rels(z, "xl/workbook.xml")
    out = {}
    for sh in ET.fromstring(z.read("xl/workbook.xml")).find("m:sheets", NS):
        part = wr[sh.get(f"{{{NS['r']}}}id")][1]
        imgs = []
        for typ, dpart in _rels(z, part).values():
            if typ != "drawing":
                continue
            dr = _rels(z, dpart)
            for anc in ET.fromstring(z.read(dpart)):
                frm = anc.find("xdr:from", NS)
                row = int(frm.find("xdr:row", NS).text) + 1 if frm is not None else None
                col = L(int(frm.find("xdr:col", NS).text) + 1) if frm is not None else None
                for blip in anc.iter(f"{{{NS['a']}}}blip"):
                    rid = blip.get(f"{{{NS['r']}}}embed")
                    if rid in dr:
                        imgs.append({"media": dr[rid][1], "row": row, "col": col})
        out[sh.get("name")] = {"state": sh.get("state", "visible"), "images": imgs}
    return out


def book_id(rel):
    stem = os.path.splitext(os.path.basename(rel))[0]
    m = re.search(r"\d{4}年\d{1,2}月数据", rel)          # 2023年X月数据\ 里的同名副本
    return stem + (f"@{m.group(0)}" if m else "")


def parse_rent_book(rel):
    path = os.path.join(SRC, rel)
    names, grids = read_grids(path)
    media = sheet_media(path)
    zone, y, mo = RENT.match(os.path.basename(rel)).groups()
    st = os.stat(path)
    sheets = []
    for name in names:
        notices = parse_rent_sheet(grids[name])
        for n in notices:
            if n["ym"] is None:
                n["ym"] = f"{y}-{int(mo):02d}"
                n["ymFrom"] = "book"            # 标题无年月(「缴费通知单」):按册月
        imgs = media.get(name, {}).get("images", [])
        for im in imgs:                          # 图挂到锚点所在(或其上方最近)的通知单
            im["notice"] = max((i for i, n in enumerate(notices) if n["titleRow"] <= (im["row"] or 0)), default=None)
        sheets.append({"name": name, "state": media.get(name, {}).get("state"), "cells": len(grids[name]),
                       "images": imgs, "notices": notices})
    with open(path, "rb") as f:
        md5 = hashlib.md5(f.read()).hexdigest()
    return {"file": rel, "id": book_id(rel), "zone": zone, "ym": f"{y}-{int(mo):02d}", "size": st.st_size,
            "mtime": datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"), "md5": md5,
            "sheetCount": len(names), "sheets": sheets}


def norm_name(t):
    return re.sub(r"\s", "", t).replace("（", "(").replace("）", ")")


def fingerprint(b):
    """同 zone 同月判重:sheet 集合(空白 sheet 不算)+ 每张通知单(抬头, 合计)。"""
    return (frozenset(sh["name"] for sh in b["sheets"] if sh["cells"] or sh["images"]),
            frozenset((sh["name"], i, n["addressee"], n["total"]) for sh in b["sheets"]
                      for i, n in enumerate(sh["notices"])))


def fp_diff(a, b):
    sa, na = fingerprint(a)
    sb, nb = fingerprint(b)
    d = []
    if sa - sb:
        d.append(f"仅 {a['id']}: {sorted(sa - sb)}")
    if sb - sa:
        d.append(f"仅 {b['id']}: {sorted(sb - sa)}")
    ka = {(x[0], x[1]): x for x in na}
    kb = {(x[0], x[1]): x for x in nb}
    for k in sorted(set(ka) & set(kb)):
        if ka[k] != kb[k]:
            d.append(f"{k[0]}#{k[1]} 合计 {ka[k][3]} vs {kb[k][3]}" +
                     (f" 抬头 {ka[k][2]} vs {kb[k][2]}" if ka[k][2] != kb[k][2] else ""))
    return d


def export_images(b, written):
    """written: sha1 → 已写出的路径。同一张图(各月册反复嵌同一份纸约截图)只落一次盘,其余硬链接过去。"""
    z = zipfile.ZipFile(os.path.join(SRC, b["file"]))
    d = os.path.join(OUT, "images", b["id"])
    os.makedirs(d, exist_ok=True)
    for sh in b["sheets"]:
        for i, im in enumerate(sh["images"], 1):
            data = z.read(im["media"])
            ext = os.path.splitext(im["media"])[1].lower()
            fn = f"{re.sub(r'[\\/:*?\"<>|]', '_', sh['name'])}_{i}{'.png' if ext == '.png' else ext}"
            path, sha = os.path.join(d, fn), hashlib.sha1(data).hexdigest()
            if os.path.exists(path):
                os.remove(path)
            try:
                os.link(written[sha], path)
            except (KeyError, OSError):
                with open(path, "wb") as f:
                    f.write(data)
                written[sha] = path
            im["file"], im["sha1"] = f"images/{b['id']}/{fn}", sha


# ───────────────────────────── 汇总表 ─────────────────────────────
def table(g, must):
    """表头行(含 must 里任一字样,前 6 行内)+ 可选的一行子表头 → [{表头: 值}]。子表头列名写成「主/子」。"""
    mc = max((c for _, c in g), default=0)
    hr = next(r for r in range(1, 7) if any(s(g.get((r, c))) in must for c in range(1, mc + 1)))
    main, cur = {}, None
    for c in range(1, mc + 1):
        if g.get((hr, c)) is not None:
            cur = s(g[(hr, c)])
        main[c] = cur
    subv = [(c, v) for (r, c), v in g.items() if r == hr + 1]
    sub = {c: s(v) for c, v in subv} if len(subv) >= 2 and all(isinstance(v, str) for _, v in subv) else {}
    heads, seen = {}, set()
    for c in range(1, mc + 1):
        h = (f"{main[c]}/{sub[c]}" if main[c] else sub[c]) if c in sub else main[c] if (hr, c) in g else None
        if h is not None and h in seen:
            h = f"{h}#{L(c)}"              # 同名列(两个「合计」之类)带列字母区分
        seen.add(h)
        heads[c] = h
    first = hr + (2 if sub else 1)
    out = []
    for r in range(first, max((r for r, _ in g), default=0) + 1):
        row = {(heads[c] or L(c)): val(g[(r, c)]) for c in range(1, mc + 1) if (r, c) in g}
        if row:
            out.append({"row": r, **row})
    return {"headerRow": hr, "subHeaderRow": hr + 1 if sub else None, "rows": out}


def parse_summaries():
    res = {}
    _, grids = read_grids(os.path.join(SRC, SUMMARY), max_col=40)
    res["contract-summary-2024-03.json"] = {"file": SUMMARY, "sheet": "汇总", **table(grids["汇总"], {"期"})}
    lease = sorted(glob.glob(os.path.join(SRC, LEASE_GLOB)),
                   key=lambda p: tuple(int(x) for x in re.findall(r"\d+", os.path.basename(p))))[-1]
    rel = os.path.relpath(lease, SRC)
    date = "-".join(f"{int(x):02d}" for x in re.findall(r"\d+", os.path.basename(lease)))
    names, grids = read_grids(lease, max_col=40)
    sheets = {}
    for n in names:
        try:
            sheets[n] = table(grids[n], {"公司", "项目", "楼号", "位置", "租户"})
        except StopIteration:
            sheets[n] = {"headerRow": None, "rows": [{"row": r, **{L(c): val(v) for (rr, c), v in sorted(grids[n].items())
                                                                  if rr == r}} for r in sorted({r for r, _ in grids[n]})]}
    res[f"lease-status-{date}.json"] = {"file": rel, "sheets": sheets}
    return res


# ───────────────────────────── 水电费册 路灯/绿化水公摊面积 ─────────────────────────────
SHARE = {"路灯公摊": "light", "路灯分摊": "light", "绿化水公摊": "green", "绿化水分摊": "green"}


def parse_utility_book(args):
    ym, zone, rel = args
    names, grids = read_grids(os.path.join(SRC, rel), max_col=16)
    out = []
    for name in names:
        g = grids[name]
        titles = sorted(r for (r, c), v in g.items() if c == 2 and isinstance(v, str) and is_title(s(v)))
        max_r = max((r for r, _ in g), default=0)
        for i, tr in enumerate(titles):
            end = titles[i + 1] if i + 1 < len(titles) else max_r + 1
            stop = next((r for (r, c), v in sorted(g.items()) if tr < r < end and s(v) == "签收日期"), end)
            cells = sorted((r, c, v) for (r, c), v in g.items() if tr <= r < stop)
            tenant = next((re.sub(r"^租户名称[:：]", "", s(v)) for r, c, v in cells if s(v).startswith("租户名称")), None)
            if tenant is None:
                continue
            period = next((re.sub(r"^计费期限[:：]", "", s(v)) for r, c, v in cells if s(v).startswith("计费期限")), None)
            pos = None
            for r, c, v in cells:
                if s(v).startswith("位置"):
                    pos = re.sub(r"^位置[:：]", "", str(v)).strip() or val(g.get((r, c + 1)))
                    break
            rec = {"ym": ym, "zone": zone, "file": rel, "sheet": name, "titleRow": tr, "title": str(g[(tr, 2)]).strip(),
                   "period": period, "tenant": tenant, "position": pos, "light": [], "green": []}
            hdr, block = {}, None
            for r, c, v in cells:
                t = s(v)
                if t == "金额":         # 表头行:整行记下(宿舍段表头多「宿舍|租赁面积」两列)
                    hdr = {s(x): cc for (rr, cc), x in g.items() if rr == r and isinstance(x, str)}
                if c == 2 and t and t != "项目" and not isinstance(v, (int, float)):
                    block = str(v).strip()
                if t in SHARE and c <= 6:
                    amt_c = hdr.get("金额")
                    area_c = next((cc for k, cc in hdr.items() if k.startswith("租赁面积")), None)
                    # 费目格与金额格之间的数:最后一个是单价(普通段「单价」列,宿舍段「用水/用电维护费」列),
                    # 普通段第一个是面积(列不固定:D/E/F/G 都有,跟着「上月行至」那列走)
                    nums = [cc for cc in range(c + 1, amt_c or c + 1) if D(g.get((r, cc))) is not None]
                    x = {"row": r, "price": val(g[(r, nums[-1])]) if nums else None,
                         "amount": val(g.get((r, amt_c))) if amt_c else None, "pool": val(g.get((r, 1))), "block": block}
                    if area_c:         # 宿舍段:公摊行不写面积,面积在上方房间行的「租赁面积」列(中间可能夹一行新水表)
                        ar = next((rr for rr in range(r - 1, r - 4, -1) if D(g.get((rr, area_c))) is not None), None)
                        x.update(areaRow=ar, areaCol=L(area_c), area=val(g.get((ar, area_c))),
                                 room=val(g.get((ar, hdr.get("宿舍", 3)))))
                    else:
                        ac = nums[0] if len(nums) >= 2 else None
                        x.update(areaRow=r, areaCol=ac and L(ac), area=val(g.get((r, ac))))
                    rec[SHARE[t]].append(x)
            out.append(rec)
    return out


# ───────────────────────────── 主流程 ─────────────────────────────
def dump(name, doc):
    with open(os.path.join(OUT, name), "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    os.makedirs(OUT, exist_ok=True)
    found = sorted(os.path.relpath(p, SRC) for p in glob.glob(os.path.join(SRC, "**", "*租金*.xlsx"), recursive=True)
                   if not os.path.basename(p).startswith("~$"))
    rent = [p for p in found if RENT.match(os.path.basename(p))]
    skipped = [{"file": p, "reason": "汇总表,另抽(contract-summary / lease-status)" if "汇总" in p or "情况表" in p
                else "不是 xlsx 包(旧 xls 改了扩展名),未读" if not zipfile.is_zipfile(os.path.join(SRC, p))
                else "不是按户通知单册(实收/损益)"} for p in found if p not in rent]
    util = [(ym, z, m[k]) for ym, m in PX.MONTHS.items() for z, k in (("一期", "p1"), ("二期", "p2"))]
    with ProcessPoolExecutor(8) as ex:
        uf = ex.map(parse_utility_book, util)
        sf = ex.submit(parse_summaries)
        books = list(ex.map(parse_rent_book, rent))
        utility = [x for part in uf for x in part]
        summaries = sf.result()

    # 判重:同 zone 同月,规范份 = 不在「X月数据」副本目录、不带 (version)的那本
    notes, groups = [], {}
    for b in books:
        groups.setdefault((b["zone"], b["ym"]), []).append(b)
    for grp in groups.values():
        grp.sort(key=lambda b: ("@" in b["id"], "version" in b["id"], b["file"]))
        canon = grp[0]
        for b in grp:
            b["duplicateOf"], b["variantOf"], b["diff"] = "", "", []
        for b in grp[1:]:
            d = fp_diff(canon, b)
            if d:
                b["variantOf"], b["diff"] = canon["file"], d
                notes.append(f"{b['id']} 与 {canon['id']} 同月不同:{'; '.join(d[:6])}{' …' if len(d) > 6 else ''}")
            else:
                b["duplicateOf"] = canon["file"]

    fails, flat, tenants, written, ym_off = [], [], {}, {}, []
    for b in books:
        if not b["duplicateOf"]:
            export_images(b, written)
        for sh in b["sheets"]:
            for i, n in enumerate(sh["notices"]):
                if not n["ok"]:
                    fails.append({"book": b["id"], "sheet": sh["name"], "notice": i, "titleRow": n["titleRow"],
                                  "total": n["total"], "sumAmount": n["sumAmount"], "problem": n.get("problem")})
                if not b["duplicateOf"]:
                    flat.append({"book": b["id"], "zone": b["zone"], "bookYm": b["ym"], "sheet": sh["name"],
                                 "sheetState": sh["state"], "notice": i, "images": sum(1 for im in sh["images"]
                                                                                     if im["notice"] == i), **n})
                    if n["addressee"]:
                        tenants.setdefault(b["zone"], set()).add(norm_name(n["addressee"]))
                    if n["ym"] != b["ym"]:
                        ym_off.append(f"{b['id']}/{sh['name']}#{i} 「{n['title']}」 合计 {n['total']}")
        dump(f"{b['id']}.json", b)
    for k, v in summaries.items():
        dump(k, v)
    dump("utility-area.json", utility)
    doc = {"generated": datetime.now().isoformat(timespec="seconds"), "src": SRC, "skipped": skipped,
           "books": [{k: b[k] for k in ("file", "id", "zone", "ym", "sheetCount", "size", "mtime", "md5",
                                        "duplicateOf", "variantOf", "diff")}
                     | {"notices": sum(len(sh["notices"]) for sh in b["sheets"]),
                        "images": sum(len(sh["images"]) for sh in b["sheets"])} for b in books],
           "tenantsByZone": {z: len(v) for z, v in sorted(tenants.items())},
           "tenantNames": {z: sorted(v) for z, v in sorted(tenants.items())},
           "notes": notes, "titleYmNotBookYm": ym_off, "uniqueImages": len(written),
           "selfcheck": {"ok": not fails, "failures": fails}, "notices": flat}
    dump("all.json", doc)
    for b in doc["books"]:
        print(f"{b['id']}: {b['sheetCount']} sheet / {b['notices']} 单 / {b['images']} 图"
              + (f" = {b['duplicateOf']}" if b["duplicateOf"] else "") + (f" ≈ {b['variantOf']}" if b["variantOf"] else ""))
    print("户数:", doc["tenantsByZone"], "| 水电公摊单:", len(utility))
    for n in notes:
        print("  注意:", n)
    print(f"自检:合计≠Σ应收 {len(fails)} 张")
    for x in fails:
        print("  ", x)
    sys.exit(0 if not fails else 1)


if __name__ == "__main__":
    main()
