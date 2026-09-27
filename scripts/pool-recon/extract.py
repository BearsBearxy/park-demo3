# -*- coding: utf-8 -*-
"""extract.py — 源册 xlsx → 当月公共电核算期望值(逐池 / 逐户 / 损耗单元),带合计自检。

    python scripts/pool-recon/extract.py 2024-02          → scripts/pool-recon/out/2024-02/expected.json

版式按内容认,不按行号:列靠表头文字定位(「本月用电量」「分摊标准」「应分摊金额」…),
池靠公式认(一期:「分摊标准」列有公式的行是池头,公式里引用的 S 行就是池的表行;
二期:「表值」列为「总」的行是块头,到下一块头为止是一个池)。
所有期望数字都取自单元格缓存值(data_only=True);公式只用来认结构(另开 data_only=False)。

自检(不合即退出码 1,expected.json 仍写出但 selfcheck.ok=false):
  一期  每块池应分摊Σ = 块合计行「应分摊金额」;每块用量Σ = 块合计行「本月用电量」
  二期  池应分摊Σ = 「二期园区公共电合计」行 W;户表每一列Σ = 户表「合计」行同列
"""
import json
import os
import re
import sys
import warnings
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")   # openpyxl 对坏打印区域的 UserWarning,和数无关

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = r"C:\financial_dashboard\2025全年发生额、预算对比"
M3 = r"2024年\2024年3月费用数据\2024年3月费用数据"

# 每月规范份(同月多份时其余只作旁证,不用的理由写在 variants)。
MONTHS = {
    "2023-08": {"p1": r"2023年\一期\水电费\一期2023年08月水电费.xlsx",
                "p2": r"2023年\二期\二期2023年08月水电费.xlsx",
                "variants": ["2023年9月数据\\一期费用\\ 同名副本:公共电分摊明细合计全同(md5 不同),旁证",
                             "2023年9月数据\\二期费用\\ 同名副本(mtime 2023-09-28 早于规范份 10-13):X 已分摊 12468.97≠12568.43,不用"]},
    "2023-09": {"p1": r"2023年\一期\水电费\一期2023年09月水电费.xlsx",
                "p2": r"2023年\二期\二期2023年09月水电费.xlsx",
                "variants": ["※一期2023年09月水电费(第一次修改模版.未完成):AE11/AE84=#REF!,未完成稿,不用",
                             "2023年10月数据\\ 下一期/二期副本:合计与规范份一致,旁证"]},
    "2023-10": {"p1": r"2023年\一期\水电费\一期2023年10月水电费.xlsx",
                "p2": r"2023年\二期\二期2023年10月水电费.xlsx",
                "variants": ["★一期2023年10月水电费-发翔海:r2 标题仍写 9 月、旧行位、AE31=#VALUE!,只可借抄表数,不用",
                             "一期(计算公共电占比变化):园区 S11=12254.18≠5005.01,占比试算稿,不用",
                             "一期 2023年11月数据\\ 副本(mtime 2026-08-11):AE11/AE62 微差,用规范份",
                             "二期(复制抄表数据至正确模版):合计 r125,W=22118.07≠20914.83,中间稿,不用",
                             "二期(计算公共电占比变化):L=40223.52 W=24866.29,试算,不用"]},
    "2023-11": {"p1": r"2023年\一期\水电费\一期2023年11月水电费.xlsx",
                "p2": r"2023年\二期\二期2023年11月水电费.xlsx",
                "variants": ["二期 2023年12月数据\\ 副本(mtime 2023-12-28):L/W 与规范份(mtime 2024-01-31,后改)不同,不用"]},
    "2023-12": {"p1": r"2023年\一期\水电费\一期2023年12月水电费.xlsx",
                "p2": r"2023年\二期\二期2023年12月水电费.xlsx",
                "variants": ["一期电费总表页名仍叫「2023年11月电费总表」,按 B1 标题认月"]},
    "2024-01": {"p1": r"2023年\一期\水电费\一期2024年1月水电费.xlsx",
                "p2": r"2023年\二期\二期2024年1月水电费.xlsx",
                "variants": ["一期电费总表页名仍叫「2023年11月电费总表」;二期公共电数据 B2 标题写 2023年12月,按文件名认月"]},
    "2024-02": {"p1": M3 + r"\一期\一期2024年2月水电费.xlsx",
                "p2": M3 + r"\二期\二期2024年2月水电费.xlsx",
                "variants": ["二期 (version 2).xlsb.xlsx(mtime 2026-08-14):合计同规范份,旁证"]},
}

# 只有抄表页、没有公摊计算页的月份:只核读数与用量(recon.py readings)
READINGS_ONLY = {
    "2024-05": M3 + r"\202405水电表数据表.xlsx",
}

TOL = Decimal("0.005")


def D(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = str(v).strip().replace(",", "")
    try:
        return Decimal(s)
    except Exception:
        return None


def num(v):
    """JSON 出口:Decimal → str(保精度,比较端再转 Decimal);非数 → None。"""
    d = D(v)
    # 浮点缓存值的尾巴(3323.5499999999997)在 8 位小数处收掉;电价最多 8 位小数,不丢信息
    return None if d is None else format(d.quantize(Decimal("1e-8")).normalize(), "f")


def ym_of(text):
    m = re.search(r"(\d{4})年(\d{1,2})月?", str(text or ""))   # 二期户表标题写「2024年2二期…」漏了「月」
    return f"{m.group(1)}-{int(m.group(2)):02d}" if m else None


class Grid:
    """一张 sheet 的缓存值 + 公式两份网格(read_only 逐行读,只解析用到的 sheet)。"""

    def __init__(self, path, name, max_col=45):
        self.name = name
        self.v, self.f = {}, {}
        for data_only, dst in ((True, self.v), (False, self.f)):
            ws = openpyxl.load_workbook(path, read_only=True, data_only=data_only)[name]
            for r, row in enumerate(ws.iter_rows(max_col=max_col, values_only=True), 1):
                for c, x in enumerate(row, 1):
                    if x is not None:
                        dst[(r, c)] = x
        self.max_row = max((r for r, _ in self.v), default=0)

    def val(self, r, c):
        return self.v.get((r, c))

    def fml(self, r, c):
        x = self.f.get((r, c))
        return x if isinstance(x, str) and x.startswith("=") else None

    def text(self, r, c):
        x = self.v.get((r, c))
        return str(x).strip() if x is not None else ""

    def header(self, must, rows=range(1, 9)):
        """找同时含 must 全部标签的表头行 → (行号, {标签: 列号})。标签按「去空白后以其开头」匹配。"""
        for r in rows:
            labels = {c: re.sub(r"\s", "", str(self.v[(r, c)])) for (rr, c) in self.v if rr == r
                      and isinstance(self.v[(r, c)], str)}
            if all(any(t.startswith(m) for t in labels.values()) for m in must):
                return r, labels
        raise ValueError(f"{self.name}: 找不到含 {must} 的表头行")


def col_of(labels, *names, required=True):
    for n in names:
        for c, t in sorted(labels.items()):
            if t.startswith(n):
                return c
    if required:
        raise ValueError(f"表头缺列 {names}")
    return None


class Book:
    def __init__(self, rel):
        self.path = os.path.join(SRC, rel)
        self.rel = rel
        self.names = openpyxl.load_workbook(self.path, read_only=True).sheetnames
        self.cache = {}

    def grid(self, pattern, exclude=None):
        hits = [n for n in self.names if re.search(pattern, n) and not (exclude and re.search(exclude, n))]
        if len(hits) != 1:
            raise ValueError(f"{self.rel}: sheet /{pattern}/ 命中 {hits}")
        if hits[0] not in self.cache:
            self.cache[hits[0]] = Grid(self.path, hits[0])
        return self.cache[hits[0]]

    def meta(self):
        st = os.stat(self.path)
        return {"file": self.rel, "size": st.st_size,
                "mtime": __import__("datetime").datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds")}


# ───────────────────────────── 抄表页(一期园区电 / 二期园区电) ─────────────────────────────
def park_usage(bk, pattern):
    """抄表页 A 列表名 → 本月用电量(总)。分摊页的 I/N 是 VLOOKUP 这里来的,两处不一致时引擎取的是这里。"""
    g = bk.grid(pattern)
    hr, lab = g.header(["区域", "上月行至", "本月用电量"])
    cS = col_of(lab, "本月用电量")
    out = {}
    for r in range(hr + 2, g.max_row + 1):
        name = g.text(r, 1)
        if name and name not in out and D(g.val(r, cS)) is not None:
            out[name] = num(g.val(r, cS))
    return out


def park_rows(bk, pattern):
    """抄表页逐行(只核读数的月份用):表名/表类/编码/倍率/上月·本月行至(总)/本月用电量(总+尖峰平谷),
    另按 ROUND((本月−上月)×倍率,2) 复算总量。本月行至为空的行照样列出(册上用量会是负数)。"""
    g = bk.grid(pattern)
    hr, lab = g.header(["区域", "上月行至", "本月用电量"])
    cI, cN, cS = col_of(lab, "上月行至"), col_of(lab, "本月行至"), col_of(lab, "本月用电量")
    cE, cG, cH = col_of(lab, "表类"), col_of(lab, "电表编码"), col_of(lab, "电表倍率")
    rows = []
    for r in range(hr + 2, g.max_row + 1):
        name = g.text(r, 1)
        if not name:
            continue
        prev, curr, h = D(g.val(r, cI)), D(g.val(r, cN)), D(g.val(r, cH))
        rows.append({"row": r, "name": name, "kind": g.text(r, cE) or None, "code": g.text(r, cG) or None,
                     "factor": num(h), "prev": num(prev), "curr": num(curr), "qty": num(g.val(r, cS)),
                     "seg": [num(g.val(r, cS + i)) for i in range(1, 5)],
                     "recalc": None if None in (prev, curr, h) else num(((curr - prev) * h).quantize(Decimal("0.01"), ROUND_HALF_UP))})
    return rows


# ───────────────────────────── 一期「公共电分摊明细」 ─────────────────────────────
S_REF = re.compile(r"(?<![A-Z$])\$?S\$?(\d+)")


def parse_p1(bk, ym, chk):
    book = park_usage(bk, r"^一期园区电$")
    g = bk.grid(r"^公共电分摊明细$")
    hr, lab = g.header(["本月用电量", "分摊标准", "应分摊金额"])
    cB, cZ = col_of(lab, "区域"), col_of(lab, "分摊范围")
    cH, cS = col_of(lab, "电表倍率"), col_of(lab, "本月用电量")
    cAA, cAC = col_of(lab, "分摊系数"), col_of(lab, "分摊标准")
    cAD, cAE, cAF = col_of(lab, "应分摊金额"), col_of(lab, "已分摊金额"), col_of(lab, "盈")
    cAG = col_of(lab, "备注", required=False)
    title = " ".join(g.text(r, c) for r in range(1, hr) for c in range(1, 6))
    chk.title("p1 公共电分摊明细", ym_of(title), ym)

    pools, groups, blocks = [], [], []
    block_rows = []
    for r in range(hr + 2, g.max_row + 1):
        if "合计" in g.text(r, cB):
            blocks.append((g.text(r, cB), r, block_rows))
            block_rows = []
        elif any((r, c) in g.v for c in (1, cB, cS, cZ, cAC, cAD, cAE)):
            block_rows.append(r)

    for bname, trow, rows in blocks:
        owned = {}
        bpools = []
        for r in rows:
            acf = g.fml(r, cAC)
            if not acf:
                continue
            srows = [int(x) for x in S_REF.findall(acf.split("/AA")[0].split("*AB")[0])] or [r]
            num_part = re.search(r"ROUND\((.*),\s*\d\s*\)\s*$", acf)
            expr = num_part.group(1) if num_part else acf
            numer = expr.split("/AA")[0].split("*AB")[0]
            consts = re.findall(r"(?<![A-Z0-9.$])([+-]?\s*\d+(?:\.\d+)?)", S_REF.sub("", numer))
            extra = sum((D(x.replace(" ", "")) for x in consts), Decimal(0))
            scale = re.search(r",\s*(\d)\s*\)\s*$", acf)
            p = {"zone": "p1", "block": bname, "headRow": r, "rows": srows,
                 "key": g.text(r, 1) or None, "members": g.text(r, cZ) or None,
                 "acFormula": acf, "roundScale": int(scale.group(1)) if scale else None,
                 "extraQty": num(extra) if extra else None,
                 "base": num(g.val(r, cAA)) if "/AA" in acf else None,
                 "std": num(g.val(r, cAC)), "note": g.text(r, cAG) if cAG else None}
            qty, cost, no_ad, meters = Decimal(0), Decimal(0), [], []
            for sr in srows:
                owned[sr] = p
                s = D(g.val(sr, cS))
                # 用量不走行至差的行(招商中心 S8 = 一期园区电!X50 + N8,N8 是写在「本月行至」位的扣度常数)
                sf = g.fml(sr, cS) or ""
                for nr in re.findall(r"\+\s*N(\d+)\b", sf):
                    extra += D(g.val(int(nr), col_of(lab, "本月行至"))) or 0
                meters.append({"row": sr, "name": g.text(sr, 1) or None, "factor": num(g.val(sr, cH)),
                               "prev": num(g.val(sr, col_of(lab, "上月行至"))),
                               "curr": num(g.val(sr, col_of(lab, "本月行至"))),
                               "qty": num(s), "bookQty": book.get(g.text(sr, 1)), "sFormula": sf or None})
                qty += s or 0
                ad = D(g.val(sr, cAD))
                if ad is None:
                    no_ad.append(sr)
                else:
                    cost += ad
            p.update(meters=meters, qty=num(qty), cost=num(cost), rowsWithoutAD=no_ad,
                     extraQty=num(extra) if extra else None)
            bpools.append(p)
        # 无「分摊标准」公式、又不属于任何池的行 = 无表行 / 手输行(联塑固定额、C 座一楼三户)
        for r in rows:
            if r in owned or not (g.text(r, 1) or g.text(r, cZ) or D(g.val(r, cAE)) is not None):
                continue
            p = {"zone": "p1", "block": bname, "headRow": r, "rows": [r], "manual": True,
                 "key": g.text(r, 1) or None, "members": g.text(r, cZ) or None,
                 "acFormula": None, "roundScale": None, "extraQty": None, "base": None,
                 "std": num(g.val(r, cAC)), "note": g.text(r, cAG) if cAG else None,
                 "meters": [], "qty": num(g.val(r, cS)), "cost": num(g.val(r, cAD)), "rowsWithoutAD": []}
            owned[r] = p
            bpools.append(p)
        # 已分摊组:AF = AE{r} − AD{a} − AD{b}…;AE 是实收(从电费总表拉回),与引擎「摊出」同义不同口径
        sum_ae = Decimal(0)
        for r in rows:
            ae = D(g.val(r, cAE))
            aff = g.fml(r, cAF) or ""
            ad_rows = [int(x) for x in re.findall(r"AD(\d+)", aff)]
            if ae is None and not ad_rows:
                continue
            groups.append({"zone": "p1", "aeRow": r, "rows": ad_rows or [r], "ae": num(ae),
                           "aeFormula": g.fml(r, cAE)})
            sum_ae += ae or 0
        if D(g.val(trow, cAE)) is not None:
            chk.eq(f"p1 {bname} 已分摊组Σ vs 合计行 AE(r{trow})", sum_ae, g.val(trow, cAE))
        # ── 自检:块内池应分摊Σ / 用量Σ = 合计行 ──
        tot_ad, tot_s = D(g.val(trow, cAD)), D(g.val(trow, cS))
        sum_cost = sum((D(p["cost"]) or 0 for p in bpools), Decimal(0))
        sum_qty = sum((D(p["qty"]) or 0 for p in bpools), Decimal(0))
        cell_ad = sum((D(g.val(r, cAD)) or 0 for r in rows), Decimal(0))
        if tot_ad is not None:
            chk.eq(f"p1 {bname} 池应分摊Σ vs 合计行 AD(r{trow})", sum_cost, tot_ad)
        else:
            # A座及园区块合计行只有 S 与 AC=ROUND(ΣS×价),无 AD 合计 → 用逐格 AD 之和核池覆盖完整
            chk.eq(f"p1 {bname} 池应分摊Σ vs 块内 AD 逐格Σ(合计行无 AD)", sum_cost, cell_ad)
        if tot_s is not None:
            chk.eq(f"p1 {bname} 池用量Σ vs 合计行 S(r{trow})", sum_qty, tot_s)
        pools.extend(bpools)
    price = next((D(g.val(p["headRow"], col_of(lab, "分摊单价"))) for p in pools
                  if D(g.val(p["headRow"], col_of(lab, "分摊单价"))) is not None), None)
    cols = {k: get_column_letter(c) for k, c in (('qty', cS), ('base', cAA), ('std', cAC), ('cost', cAD), ('alloc', cAE))}
    return pools, groups, price, cols


# ───────────────────────────── 二期「公共电数据」 ─────────────────────────────
def eff_base(formula, t):
    """二期 V 公式里跟在 /T{r} 后面的 /k*m 折算 → 等效基数(四车间电梯 /T/4*3 → T×4/3;广联 /T/4 → T×4)。"""
    t = D(t)
    m = re.search(r"/T\$?\d+((?:\s*[/*]\s*\d+(?:\.\d+)?)*)", formula or "")
    if t is None or not m:
        return t
    for op, k in re.findall(r"([/*])\s*(\d+(?:\.\d+)?)", m.group(1)):
        t = t * D(k) if op == "/" else t / D(k)
    return t


def parse_p2(bk, ym, chk):
    book = park_usage(bk, r"^二期园区电$")
    park = bk.grid(r"^二期园区电$")
    g = bk.grid(r"^公共电数据$")
    hr, lab = g.header(["本月用电量", "分摊标准", "应分摊金额"])
    cB, cD = col_of(lab, "区域"), col_of(lab, "企业名称")
    cH, cI, cL = col_of(lab, "电表倍率"), col_of(lab, "表值"), col_of(lab, "本月用电量")
    cS, cT, cU = col_of(lab, "分摊范围"), col_of(lab, "分摊系数"), col_of(lab, "分摊单价")
    cV, cW, cX = col_of(lab, "分摊标准"), col_of(lab, "应分摊金额"), col_of(lab, "已分摊金额")
    cZ = col_of(lab, "备注", required=False)
    title = " ".join(g.text(r, c) for r in range(1, hr) for c in range(1, 6))
    chk.title("p2 公共电数据", ym_of(title), ym)

    trow = next(r for r in range(hr + 1, g.max_row + 1) if "合计" in g.text(r, cB))
    heads = [r for r in range(hr + 1, trow) if g.text(r, cI) == "总" and g.text(r, 1)]
    pools, seg_price = [], {}
    for i, h in enumerate(heads):
        end = heads[i + 1] if i + 1 < len(heads) else trow
        seg = {"sharp": Decimal(0), "peak": None, "flat": None, "valley": None}
        gl_row, sharp_row, ratio = None, None, None
        for r in range(h + 1, end):
            lab_r, lv, lf = g.text(r, cI), D(g.val(r, cL)), g.fml(r, cL) or ""
            # 隐藏的尖价折出行 = 紧跟「尖峰」行的无标签行(一般是 ×二期园区电!AB4,也见过 AB10 与手输数)
            hidden = not lab_r and sharp_row is not None and r == sharp_row + 1
            if lab_r == "尖峰":
                sharp_row = r
            if lab_r == "尖峰" or hidden:
                seg["sharp"] += lv or 0          # 尖峰行 + 隐藏的尖价折出行(比率×总量)
            elif lab_r in ("峰", "平", "谷"):
                seg[{"峰": "peak", "平": "flat", "谷": "valley"}[lab_r]] = lv
            if hidden:
                ref = re.search(r"园区电!\$?(AB)\$?(\d+)", lf)
                if ref:
                    ratio = {"cell": f"二期园区电!AB{ref.group(2)}", "value": num(park.val(int(ref.group(2)), 28))}
                else:                             # 手输:按 隐藏行 / (尖峰行 + 隐藏行) 反推
                    tot = (D(g.val(sharp_row, cL)) or 0) + (lv or 0)
                    ratio = {"cell": f"公共电数据!L{r} 手输", "value": num((lv or 0) / tot) if tot else None}
            # 段单价(含 0.16 维护费):尖价在隐藏折出行,峰/平/谷在各自行;取第一块齐全的
            k = "sharp" if hidden else {"峰": "peak", "平": "flat", "谷": "valley"}.get(lab_r)
            if k and k not in seg_price and D(g.val(r, cU)) is not None:
                seg_price[k] = num(g.val(r, cU))
            if g.text(r, cU) == "广联分摊":
                gl_row = r
        vf = g.fml(h, cV) or ""
        scale = re.search(r",\s*(\d)\s*\)", vf)
        x = g.val(h, cX)
        p = {"zone": "p2", "block": g.text(h, cB) or None, "headRow": h, "rows": list(range(h, end)),
             "key": g.text(h, 1), "name": f"{g.text(h, cB)}{g.text(h, cD)}", "members": g.text(h, cS) or None,
             "factor": num(g.val(h, cH)), "qty": num(g.val(h, cL)), "bookQty": book.get(g.text(h, 1)),
             "qtySharp": num(seg["sharp"]), "qtyPeak": num(seg["peak"]), "qtyFlat": num(seg["flat"]),
             "qtyValley": num(seg["valley"]), "sharpRatio": ratio,
             "cost": num(g.val(h, cW)), "std": num(g.val(h, cV)),
             "base": num(eff_base(vf, g.val(h, cT))), "baseCell": num(g.val(h, cT)),
             "vFormula": vf or None, "roundScale": int(scale.group(1)) if scale else None,
             "allocated": num(x), "allocatedText": None if D(x) is not None else (str(x) if x is not None else None),
             "note": g.text(h, cZ) if cZ else None}
        pools.append(p)
        if gl_row:
            pools.append({"zone": "p2", "block": p["block"], "headRow": gl_row, "rows": [gl_row],
                          "key": p["key"] + "#广联", "name": "广联分摊", "members": "广联分摊",
                          "refOf": p["key"], "qty": p["qty"], "cost": None, "std": num(g.val(gl_row, cV)),
                          "base": num(eff_base(g.fml(gl_row, cV), g.val(h, cT))),
                          "vFormula": g.fml(gl_row, cV), "allocated": None})
    tot_w = D(g.val(trow, cW))
    sum_w = sum((D(p["cost"]) or 0 for p in pools), Decimal(0))
    cell_w = sum((D(g.val(r, cW)) or 0 for r in range(hr + 1, trow)), Decimal(0))
    chk.eq(f"p2 池应分摊Σ vs 合计行 W(r{trow})", sum_w, tot_w)
    chk.eq(f"p2 池应分摊Σ vs W 列逐格Σ(查块外散落金额)", sum_w, cell_w)
    totals = {"row": trow, "qty": num(g.val(trow, cL)), "cost": num(tot_w), "segPrice": seg_price,
              "cols": {k: get_column_letter(c) for k, c in (("qty", cL), ("base", cT), ("std", cV), ("cost", cW),
                                                            ("alloc", cX))},
              "allocated": num(g.val(trow, cX)), "headQtySum": num(sum((D(p["qty"]) or 0 for p in pools), Decimal(0)))}
    return pools, totals


# ───────────────────────────── 损耗表 ─────────────────────────────
def parse_loss(bk, zone, ym, chk):
    g = bk.grid(r"园区损耗$")
    hr, lab = g.header(["收取租户损耗率"])
    cB = col_of(lab, "位置")
    cols = {"c": col_of(lab, "变压器用电量", "总表用电量"), "cable": col_of(lab, "铝缆用电量", required=False),
            "d": col_of(lab, "分表用电量"), "e": col_of(lab, "损耗量"), "rawRate": col_of(lab, "原损耗率"),
            "rate": col_of(lab, "收取租户损耗率"), "adjRate": col_of(lab, "调整损耗", required=False)}
    if zone == "p1":
        cols["g"] = col_of(lab, "分摊用电度数")
    else:
        cols["adjQty"] = col_of(lab, "分摊/调整用电度数", "分摊用电度数")
    title = " ".join(g.text(r, c) for r in range(1, hr) for c in range(1, 6))
    chk.title(f"{zone} 损耗表", ym_of(title), ym)
    park = bk.grid(r"^一期园区电$" if zone == "p1" else r"^二期园区电$")
    _, plab = park.header(["区域", "上月行至", "本月用电量"])
    pS = col_of(plab, "本月用电量")

    def members(dfml):
        """D = 抄表页!S{合计行} → 合计行 SUM(S a:S b) − S x … + X y → 段落成员(下钻用)。"""
        m = re.search(r"园区电!\$?S\$?(\d+)", dfml or "")
        if not m:
            return None
        tf = park.fml(int(m.group(1)), pS) or ""
        rng = re.search(r"SUM\(S(\d+):S(\d+)\)", tf)
        if not rng:
            return None
        minus = {int(x) for x in re.findall(r"-\s*S(\d+)", tf)}
        rows = [r for r in range(int(rng.group(1)), int(rng.group(2)) + 1) if r not in minus]
        out = [{"row": r, "name": park.text(r, 1), "qty": num(park.val(r, pS))} for r in rows
               if park.text(r, 1) or D(park.val(r, pS))]
        for x in re.findall(r"\+\s*([A-Z]+\d+)", tf):       # A座 +X50 招商净额之类的非表项
            out.append({"row": None, "name": f"{park.name}!{x}", "qty": None})
        return out
    units = []
    for r in range(hr + 1, g.max_row + 1):
        pos = g.text(r, cB)
        if not pos or "合计" in pos:
            if "合计" in pos:
                break
            continue
        u = {"zone": zone, "row": r, "label": pos}
        for k, c in cols.items():
            if c:
                u[k] = num(g.val(r, c))
        u["rateFormula"] = g.fml(r, cols["rate"])
        u["dMembers"] = members(g.fml(r, cols["d"]))
        units.append(u)
    return units


# ───────────────────────────── 户表(逐户 × 费目) ─────────────────────────────
CATS = {"消防用电": "fire", "电梯用电": "elevator", "路灯公摊": "light", "线路损耗": "loss", "绿化水公摊": "green"}


def parse_tenants(g, zone, ym, chk, skip_groups=()):
    hr, lab = g.header(["企业名称"])
    sub = {c: re.sub(r"\s", "", g.text(hr + 1, c)) for c in range(1, 45) if g.text(hr + 1, c)}
    starts = sorted((c, t) for c, t in lab.items() if c > 2)
    title = " ".join(g.text(r, c) for r in range(1, hr) for c in range(1, 6))
    chk.title(f"{zone} {g.name}", ym_of(title), ym)

    def group_of(c):
        return next((t for s, t in reversed(starts) if s <= c), "")
    cat_cols = [(c, CATS[t], group_of(c)) for c, t in sorted(sub.items()) if t in CATS]
    trow = next(r for r in range(hr + 2, g.max_row + 1) if g.text(r, 2) == "合计")
    out, cur, taken = {}, None, defaultdict(Decimal)   # taken:每列真正记到户上的额(自检用它,不用整列)
    for r in range(hr + 2, trow):
        if g.text(r, 2):
            cur = g.text(r, 2)
        if cur is None:
            continue
        for c, cat, grp in cat_cols:
            v = D(g.val(r, c))
            if v is None or v == 0:
                continue
            skipped = any(s in grp for s in skip_groups)
            t = out.setdefault(cur, {"zone": zone, "name": cur, "rows": [], "amounts": {}, "excluded": {}})
            if r not in t["rows"]:
                t["rows"].append(r)
            bucket = t["excluded"] if skipped else t["amounts"]
            bucket[cat] = num((D(bucket.get(cat)) or 0) + v)
            taken[c] += v
    for c, cat, grp in cat_cols:
        tot = D(g.val(trow, c))
        s = taken[c]
        if tot is not None:
            chk.eq(f"{zone} {g.name} {grp}·{cat} 列Σ vs 合计行(r{trow})", s, tot)
    return list(out.values())


class Check:
    def __init__(self):
        self.errors, self.passed, self.warnings = [], [], []

    def eq(self, tag, got, want):
        got, want = D(got) or Decimal(0), D(want)
        if want is None or abs(got - want) > TOL:
            self.errors.append(f"{tag}: 抽取 {got} ≠ 册上 {want}")
        else:
            self.passed.append(f"{tag}: {want}")

    def title(self, tag, found, ym):
        if found != ym:
            self.warnings.append(f"{tag} 标题月 {found} ≠ 账期 {ym}(按文件认月,继续)")


def extract(ym):
    if ym not in MONTHS:
        raise SystemExit(f"未登记的月份 {ym};已登记 {sorted(MONTHS)}")
    m = MONTHS[ym]
    chk = Check()
    b1, b2 = Book(m["p1"]), Book(m["p2"])
    p1_pools, p1_groups, p1_price, p1_cols = parse_p1(b1, ym, chk)
    p2_pools, p2_totals = parse_p2(b2, ym, chk)
    loss = parse_loss(b1, "p1", ym, chk) + parse_loss(b2, "p2", ym, chk)
    tenants = (parse_tenants(b1.grid(r"^\d{4}年\d{1,2}月电费总表$"), "p1", ym, chk, skip_groups=("居民",))
               + parse_tenants(b2.grid(r"^本月用电数据统计$"), "p2", ym, chk)
               + parse_tenants(b2.grid(r"^本月用水数据统计$"), "p2", ym, chk, skip_groups=("居民",)))
    # 二期水表与电表同名户合并成一户(绿化水只在水表)
    merged = {}
    for t in tenants:
        k = (t["zone"], t["name"])
        if k in merged:
            dst = merged[k]
            dst["rows"] += t["rows"]
            for part in ("amounts", "excluded"):
                for cat, v in t[part].items():
                    dst[part][cat] = num((D(dst[part].get(cat)) or 0) + D(v))
        else:
            merged[k] = t
    return {
        "ym": ym, "sources": {"p1": b1.meta(), "p2": b2.meta()}, "variants": m["variants"],
        "pools": p1_pools + p2_pools, "p1AllocGroups": p1_groups, "p2Totals": p2_totals,
        "loss": loss, "tenants": list(merged.values()),
        "cols": {"p1": p1_cols, "p2": p2_totals["cols"]},   # 证据单元格用的列字母(按表头认出)
        "prices": {"p1Unit": num(p1_price), "p2Seg": p2_totals["segPrice"]},   # 含维护费的册上单价
        "selfcheck": {"ok": not chk.errors, "errors": chk.errors, "warnings": chk.warnings,
                      "passed": chk.passed},
    }


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ym = sys.argv[1]
    out_dir = os.path.join(HERE, "out", ym)
    os.makedirs(out_dir, exist_ok=True)
    doc = extract(ym)
    path = os.path.join(out_dir, "expected.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)
    sc = doc["selfcheck"]
    n1 = sum(1 for p in doc["pools"] if p["zone"] == "p1")
    print(f"{ym}: 一期池 {n1} / 二期池 {len(doc['pools']) - n1} / 损耗单元 {len(doc['loss'])} / 户 {len(doc['tenants'])}"
          f" | 自检 通过 {len(sc['passed'])} 项 失败 {len(sc['errors'])} 项 → {path}")
    for w in sc["warnings"]:
        print("  注意:", w)
    for e in sc["errors"]:
        print("  自检失败:", e)
    sys.exit(0 if sc["ok"] else 1)


if __name__ == "__main__":
    main()
